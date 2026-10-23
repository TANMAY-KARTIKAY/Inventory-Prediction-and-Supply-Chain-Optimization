# -*- coding: utf-8 -*-
"""
Master Builder: Compiles the complete ~8,000-word journal manuscript into both:
1. paper/manuscript.md
2. paper/manuscript.docx
"""

import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')

# Import section modules
sys.path.append(os.path.join(r"C:\Users\tanma\OneDrive\Desktop\project(1)\paper", "sections"))

from sec1_front_and_intro import TITLE, AUTHORS, AFFILIATIONS, ABSTRACT, KEYWORDS, SECTION_1_TEXT
from sec2_literature import SECTION_2_TEXT
from sec3_problem import SECTION_3_TEXT
from sec4_methodology import SECTION_4_TEXT
from sec5_safeguards import SECTION_5_TEXT
from sec6_results import SECTION_6_TEXT
from sec7_discussion import SECTION_7_TEXT
from sec8_limitations_conclusion_refs import SECTION_8_9_TEXT

PAPER_DIR = r"C:\Users\tanma\OneDrive\Desktop\project(1)\paper"
MD_PATH = os.path.join(PAPER_DIR, "manuscript.md")
DOCX_PATH = os.path.join(PAPER_DIR, "manuscript.docx")

# Color palette for academic Word document
COLOR_NAVY_HEX = "1B365D"
COLOR_SLATE_HEX = "2C5E8A"
COLOR_LIGHT_GRAY = "F4F6F9"
COLOR_BORDER = "CCCCCC"
COLOR_NAVY_RGB = RGBColor(27, 54, 93)
COLOR_SLATE_RGB = RGBColor(44, 94, 138)
COLOR_TEXT_RGB = RGBColor(34, 34, 34)
COLOR_MUTED_RGB = RGBColor(90, 90, 90)

def sanitize_xml(s):
    if not isinstance(s, str):
        return s
    return "".join(c for c in s if ord(c) >= 32 or c in "\t\n\r")

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def build_markdown():
    """Assembles the full Markdown text."""
    parts = []
    
    # Title & Front Matter
    parts.append(f"# {TITLE}\n")
    
    # Authors
    author_line = " **" + "**, **".join([f"{a['name']}^{{{a['affiliation']}}}" for a in AUTHORS]) + "**\n"
    parts.append(author_line)
    
    # Affiliations
    for idx, aff in enumerate(AFFILIATIONS, 1):
        parts.append(f"$^{idx}$ {aff}  ")
    parts.append("\n---\n")
    
    # Abstract
    parts.append("## Abstract\n")
    parts.append(ABSTRACT.strip() + "\n")
    
    # Keywords
    parts.append(f"**Keywords:** {', '.join(KEYWORDS)}\n")
    parts.append("\n---\n")
    
    # Sections
    parts.append(SECTION_1_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_2_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_3_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_4_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_5_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_6_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_7_TEXT.strip())
    parts.append("\n---\n")
    parts.append(SECTION_8_9_TEXT.strip())
    
    full_md = "\n\n".join(parts)
    
    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write(full_md)
        
    word_count = len(re.findall(r'\b\w+\b', full_md))
    print(f"Markdown written to: {MD_PATH}")
    print(f"Total Manuscript Word Count: {word_count:,} words")
    return full_md, word_count

def parse_markdown_table(lines):
    """Parses a markdown table into headers and rows."""
    header_line = lines[0]
    headers = [c.strip() for c in header_line.split('|')[1:-1]]
    
    rows = []
    for line in lines[2:]:
        if '|' in line:
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if len(cells) == len(headers):
                rows.append(cells)
    return headers, rows

def build_docx(full_md):
    """Compiles the manuscript into Microsoft Word format."""
    doc = docx.Document()
    
    # Page setup: Standard 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True
        
        # Header / Footer
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Bridging Prediction and Inventory Optimization | Kartikay et al.")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED_RGB
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(9)
        frun.font.color.rgb = COLOR_MUTED_RGB
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run(TITLE)
    r_title.font.name = "Georgia"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY_RGB
    
    # Authors
    p_auth = doc.add_paragraph()
    p_auth.paragraph_format.space_after = Pt(4)
    for i, a in enumerate(AUTHORS):
        r_name = p_auth.add_run(a['name'])
        r_name.font.name = "Calibri"
        r_name.font.size = Pt(11)
        r_name.font.bold = True
        r_sup = p_auth.add_run(f" [{a['affiliation']}]")
        r_sup.font.name = "Calibri"
        r_sup.font.size = Pt(9)
        r_sup.font.superscript = True
        if i < len(AUTHORS) - 1:
            p_auth.add_run(", ")
            
    # Affiliations
    for idx, aff in enumerate(AFFILIATIONS, 1):
        p_aff = doc.add_paragraph()
        p_aff.paragraph_format.space_after = Pt(2)
        r_idx = p_aff.add_run(f"[{idx}] ")
        r_idx.font.name = "Calibri"
        r_idx.font.size = Pt(8.5)
        r_idx.font.bold = True
        r_text = p_aff.add_run(aff)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(8.5)
        r_text.font.color.rgb = COLOR_MUTED_RGB

    # Abstract Callout Box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(14)
    p_box.paragraph_format.space_after = Pt(4)
    r_abshdr = p_box.add_run("ABSTRACT")
    r_abshdr.font.name = "Calibri"
    r_abshdr.font.size = Pt(10)
    r_abshdr.font.bold = True
    r_abshdr.font.color.rgb = COLOR_NAVY_RGB
    
    for para in ABSTRACT.strip().split('\n\n'):
        p_abs = doc.add_paragraph()
        p_abs.paragraph_format.space_after = Pt(6)
        p_abs.paragraph_format.left_indent = Inches(0.25)
        p_abs.paragraph_format.right_indent = Inches(0.25)
        p_abs.paragraph_format.line_spacing = 1.15
        r_abs = p_abs.add_run(sanitize_xml(para.strip()))
        r_abs.font.name = "Calibri"
        r_abs.font.size = Pt(9.5)
        r_abs.font.italic = True
        r_abs.font.color.rgb = COLOR_TEXT_RGB

    # Keywords
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(4)
    p_kw.paragraph_format.space_after = Pt(16)
    p_kw.paragraph_format.left_indent = Inches(0.25)
    r_kwhdr = p_kw.add_run("Keywords: ")
    r_kwhdr.font.name = "Calibri"
    r_kwhdr.font.size = Pt(9)
    r_kwhdr.font.bold = True
    r_kwhdr.font.color.rgb = COLOR_NAVY_RGB
    r_kw = p_kw.add_run("; ".join(KEYWORDS))
    r_kw.font.name = "Calibri"
    r_kw.font.size = Pt(9)
    r_kw.font.color.rgb = COLOR_MUTED_RGB
    
    # Process text lines
    lines = full_md.split('\n')
    i = 0
    in_code = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Check code block
        if line.strip().startswith('```'):
            in_code = not in_code
            if not in_code and table_lines:
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.3)
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(6)
                r_c = p_code.add_run(sanitize_xml('\n'.join(table_lines)))
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(8)
                r_c.font.color.rgb = COLOR_SLATE_RGB
                table_lines = []
            i += 1
            continue
            
        if in_code:
            table_lines.append(line)
            i += 1
            continue
            
        # Check markdown table
        if line.strip().startswith('|') and '|' in line[1:]:
            t_block = [line]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
                i += 1
                t_block.append(lines[i])
            if len(t_block) >= 3:
                try:
                    headers, rows = parse_markdown_table(t_block)
                    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl, COLOR_BORDER)
                    
                    # Headers
                    hdr_row = tbl.rows[0]
                    for c_idx, h_text in enumerate(headers):
                        c = hdr_row.cells[c_idx]
                        set_cell_background(c, COLOR_NAVY_HEX)
                        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
                        p_c = c.paragraphs[0]
                        p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        rc = p_c.add_run(sanitize_xml(h_text.replace('*', '').replace('$', '')))
                        rc.font.name = "Calibri"
                        rc.font.size = Pt(8.5)
                        rc.font.bold = True
                        rc.font.color.rgb = RGBColor(255, 255, 255)
                        
                    # Rows
                    for r_idx, r_cells in enumerate(rows):
                        row = tbl.rows[r_idx + 1]
                        bg = COLOR_LIGHT_GRAY if r_idx % 2 == 1 else "FFFFFF"
                        for c_idx, cell_text in enumerate(r_cells):
                            c = row.cells[c_idx]
                            set_cell_background(c, bg)
                            set_cell_margins(c, top=70, bottom=70, left=120, right=120)
                            p_c = c.paragraphs[0]
                            p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            rc = p_c.add_run(sanitize_xml(cell_text.replace('*', '').replace('$', '')))
                            rc.font.name = "Calibri"
                            rc.font.size = Pt(8)
                            rc.font.color.rgb = COLOR_TEXT_RGB
                            
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_before = Pt(2)
                    p_sp.paragraph_format.space_after = Pt(6)
                except Exception as e:
                    pass
            i += 1
            continue
            
        # Check Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(line[2:].strip()))
            r.font.name = "Georgia"
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY_RGB
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(line[3:].strip()))
            r.font.name = "Calibri"
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_SLATE_RGB
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sanitize_xml(line[4:].strip()))
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = COLOR_TEXT_RGB
        elif line.strip().startswith('---'):
            pass
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            
            clean_text = line.strip()
            if clean_text.startswith('* ') or clean_text.startswith('- '):
                p.paragraph_format.left_indent = Inches(0.25)
                clean_text = "• " + clean_text[2:]
            elif re.match(r'^\d+\.\s', clean_text):
                p.paragraph_format.left_indent = Inches(0.25)
                
            r = p.add_run(sanitize_xml(clean_text))
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_TEXT_RGB
            
        i += 1
        
    doc.save(DOCX_PATH)
    print(f"Word document successfully compiled to: {DOCX_PATH}")

if __name__ == "__main__":
    full_md, wc = build_markdown()
    build_docx(full_md)
    print("\n[SUCCESS] Both academic outputs generated successfully!")
